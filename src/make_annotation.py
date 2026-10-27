#!/usr/bin/env python3
"""Build the human-annotation workbooks and the hidden answer key.

    python src/make_annotation.py

Sample (fixed random seed, so it is reproducible):
  * 202 words from REAL TurEngMix text (train + validation splits only; the test
    split is never shown to anyone), stratified by gold label so that rare labels
    have enough items: MIXED 30, AMBIGUOUS 20, OTHER 2 (all that exist), NE 40,
    EN 50, TR 60.
  * 200 words from OUR generated data (first 2,500 sentences): MIXED 45,
    AMBIGUOUS 10, OTHER 5, NE 40, EN 40, TR 60.
  * Real and generated items are shuffled together. Annotators are not told
    which is which.

Writes
  annotation/annotation_ibrahim.xlsx, annotation_mahmut.xlsx, annotation_tugra.xlsx
      (identical items; a drop-down list for the label)
  annotation/key.csv   (item, source, reference, word, reference_label)
      -> keep this file private until all three workbooks are back.

Because the sample is stratified, overall agreement here is not directly
comparable with TurEngMix's own kappa (they sampled differently); the
per-label numbers are the useful ones.
"""

from __future__ import annotations

import csv
import json
import random
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lid_io import GENERATED_PATH, REPO_ROOT, load_split_sentences  # noqa: E402

SEED = 2026
WINDOW = 12
QUOTA_REAL = {"MIXED": 30, "AMBIGUOUS": 20, "OTHER": 2, "NE": 40, "EN": 50, "TR": 60}
QUOTA_GEN = {"MIXED": 45, "AMBIGUOUS": 10, "OTHER": 5, "NE": 40, "EN": 40, "TR": 60}
ANNOTATORS = ["ibrahim", "mahmut", "tugra"]
LABELS = ["TR", "EN", "MIXED", "NE", "AMBIGUOUS", "OTHER"]
OUT = REPO_ROOT / "annotation"


def pick(pool, n, key, rng):
    """n items, preferring different sentences."""
    cands = pool[:]
    rng.shuffle(cands)
    seen, out = set(), []
    for c in cands:
        k = key(c)
        if k in seen and len(cands) > 3 * n:
            continue
        seen.add(k)
        out.append(c)
        if len(out) == n:
            break
    return out


def main() -> int:
    rng = random.Random(SEED)
    items = []

    real = load_split_sentences("train") + load_split_sentences("validation")
    pool: dict[str, list] = {}
    for s in real:
        for i, (t, lab) in enumerate(zip(s.tokens, s.labels)):
            if t.strip('.,!?;:"()[]'):
                pool.setdefault(lab, []).append((s, i))
    for lab, n in QUOTA_REAL.items():
        for s, i in pick(pool[lab], n, lambda c: (c[0].doc_id, c[0].sent_id), rng):
            items.append({"source": "real", "ref": f"{s.doc_id}/{s.sent_id}/{i}",
                          "tokens": s.tokens, "i": i, "label": s.labels[i]})

    gen = [json.loads(l) for l in open(GENERATED_PATH, encoding="utf-8")][:2500]
    pool = {}
    for r in gen:
        for i, lab in enumerate(r["labels"]):
            pool.setdefault(lab, []).append((r, i))
    for lab, n in QUOTA_GEN.items():
        for r, i in pick(pool[lab], n, lambda c: c[0]["id"], rng):
            items.append({"source": "generated", "ref": f"{r['id']}/{i}",
                          "tokens": r["tokens"], "i": i, "label": r["labels"][i]})

    rng.shuffle(items)
    for n, it in enumerate(items, start=1):
        t, i = it["tokens"], it["i"]
        lo, hi = max(0, i - WINDOW), min(len(t), i + WINDOW + 1)
        ctx = (["…"] if lo > 0 else []) + t[lo:i] + [f"[[{t[i]}]]"] + t[i + 1:hi] \
            + (["…"] if hi < len(t) else [])
        it.update(item=n, context=" ".join(ctx), word=t[i])

    OUT.mkdir(exist_ok=True)
    with open(OUT / "key.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["item", "source", "reference", "word", "reference_label"])
        for it in items:
            w.writerow([it["item"], it["source"], it["ref"], it["word"], it["label"]])

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    for name in ANNOTATORS:
        wb = Workbook()
        info = wb.active
        info.title = "Read me"
        lines = [
            ("How to label", True),
            ("Each row shows a sentence. One word is marked like [[this]]. Choose that word's label in column D.", False),
            ("Work alone. Do not discuss items with anyone. Do not use AI. Answer every row (best guess if unsure).", False),
            ("Full rules: report/annotation_guide.md in the GitHub repo. Read it first (10 minutes).", False),
            ("", False),
            ("Labels", True),
            ("TR — Turkish word (also borrowed words written the Turkish way: telefon, internet, site, film)", False),
            ("EN — English word with no Turkish ending (meeting, deadline, online)", False),
            ("MIXED — English word + Turkish ending (linki, meeting'e, download'ladım)", False),
            ("NE — any word that is part of a name, even with a Turkish ending (netflix'te, galatasaray, the office)", False),
            ("AMBIGUOUS — used in both languages and the sentence does not decide (platform, web, tv, hey) — use rarely", False),
            ("OTHER — a third language (Japanese, German, French …)", False),
            ("", False),
            ("Send the file back to Shahem when finished.", True),
        ]
        for r, (text, bold) in enumerate(lines, start=1):
            c = info.cell(row=r, column=1, value=text)
            c.font = Font(bold=bold, size=12 if bold else 11)
        info.column_dimensions["A"].width = 120

        ws = wb.create_sheet("Label here")
        header = ["item", "sentence (word to label is in [[ ]])", "word", "your label", "comment (optional)"]
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="305496")
        for it in items:
            ws.append([it["item"], it["context"], it["word"], None, None])
        dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle="Invalid label",
                            error="Choose one of: " + ", ".join(LABELS))
        ws.add_data_validation(dv)
        dv.add(f"D2:D{len(items) + 1}")
        for col, width in zip("ABCDE", (7, 95, 22, 14, 30)):
            ws.column_dimensions[col].width = width
        for row in ws.iter_rows(min_row=2):
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
            row[3].fill = PatternFill("solid", fgColor="FFF2CC")
        ws.freeze_panes = "A2"
        wb.active = 1
        wb.save(OUT / f"annotation_{name}.xlsx")

    from collections import Counter
    print(f"{len(items)} items:", dict(Counter(it["source"] for it in items)))
    print("per label:", dict(Counter((it["source"], it["label"]) for it in items)))
    print("wrote", ", ".join(f"annotation_{n}.xlsx" for n in ANNOTATORS), "and key.csv")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
